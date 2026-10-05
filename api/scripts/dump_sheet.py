import sys
import json
import openpyxl
from datetime import datetime

path = sys.argv[1]
sheet_name = sys.argv[2]
header_row = int(sys.argv[3])

wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
ws = wb[sheet_name]

headers = {}
for cell in next(ws.iter_rows(min_row=header_row, max_row=header_row)):
    if cell.value is not None and str(cell.value).strip() != "":
        headers[cell.column] = str(cell.value)

rows = []
excel_row = header_row
for row in ws.iter_rows(min_row=header_row + 1, values_only=False):
    excel_row += 1
    obj = {"_excelRow": excel_row}
    has_any = False
    for cell in row:
        name = headers.get(cell.column)
        if name is None:
            continue
        v = cell.value
        if isinstance(v, datetime):
            v = v.isoformat()
        obj[name] = v
        if v is not None and str(v).strip() != "":
            has_any = True
    if has_any:
        rows.append(obj)

wb.close()
print(json.dumps(rows))
